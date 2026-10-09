"""
core_engine.py
==============
Modular backend engine for Planogram (POG) AI Audits.
Handles parallel photo downloading, YOLO inference, and formatted Excel generation.
"""

import os
import io
import time
import zipfile
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor, as_completed
from PIL import Image
import requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from ultralytics import YOLO

BASE_DIR = Path(__file__).resolve().parent

# Hardcoded calibrated configurations - No arbitrary slider tweaking allowed
AUDIT_MODELS = {
    "Rubber Mat Audit": {
        "id": "mat",
        "name": "Rubber Mat Audit",
        "weights": BASE_DIR / "models" / "yolov8n_mat.pt",
        "conf": 0.25,
        "classes": {0: "rubber_mat_open", 1: "rubber_mat_boxed"},
        "target_desc": "Rubber Mat (Open & Boxed)"
    },
    "Threshold Fixture Audit": {
        "id": "fixture",
        "name": "Threshold Fixture Audit",
        "weights": BASE_DIR / "models" / "yolov8n_fixture.pt",
        "conf": 0.710,
        "classes": {0: "threshold_fixture"},
        "target_desc": "Threshold Display Fixture"
    }
}

CACHE_DIR = BASE_DIR / "cache" / "downloaded_images"
CACHE_DIR.mkdir(parents=True, exist_ok=True)

USER_AGENT = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"

_LOADED_MODELS = {}

def get_yolo_model(model_key: str):
    """Loads and caches the YOLO model in memory."""
    if model_key not in _LOADED_MODELS:
        weights_path = AUDIT_MODELS[model_key]["weights"]
        if not weights_path.exists():
            raise FileNotFoundError(f"Model weights not found at {weights_path}")
        _LOADED_MODELS[model_key] = YOLO(str(weights_path))
    return _LOADED_MODELS[model_key]


def identify_columns(headers):
    """Identifies photo, store/location, and bay columns from headers."""
    photo_col = -1
    store_col = -1
    bay_col = -1

    for idx, h in enumerate(headers):
        h_clean = str(h).strip().lower() if h else ""
        if h_clean == "photo link":
            photo_col = idx
        elif h_clean in ("location", "store", "store id", "store #"):
            store_col = idx
        elif h_clean in ("bay id", "bay", "bay number"):
            bay_col = idx

    # Schema fallbacks matching standard retail POG sheets
    if photo_col == -1:
        photo_col = 11 if len(headers) > 11 else len(headers) - 1
    if store_col == -1:
        store_col = 0
    if bay_col == -1:
        bay_col = 4 if len(headers) > 4 else 1

    return photo_col, store_col, bay_col


def download_all_photos_parallel(tasks, download_dir, max_workers=20, progress_callback=None):
    """
    Downloads photos in parallel with thread pooling and persistent session.
    tasks: list of (idx, row_num, store_val, photo_val, url)
    """
    session = requests.Session()
    adapter = requests.adapters.HTTPAdapter(pool_connections=max_workers, pool_maxsize=max_workers, max_retries=2)
    session.mount('https://', adapter)
    session.mount('http://', adapter)
    session.headers.update({"User-Agent": USER_AGENT})

    results = {}
    total = len(tasks)

    def _fetch_one(task):
        idx, row_num, store_val, photo_val, url = task
        if not url:
            return idx, None, "No URL in cell"

        clean_name = f"row_{row_num}_store_{store_val}_{Path(photo_val).name if photo_val else 'photo.jpg'}"
        clean_name = "".join(c if c.isalnum() or c in "._-" else "_" for c in clean_name)
        save_path = download_dir / clean_name

        # Cached check
        if save_path.exists() and save_path.stat().st_size > 1024:
            return idx, save_path, "cached"

        try:
            resp = session.get(url, timeout=20)
            if resp.status_code == 200:
                with open(save_path, "wb") as f:
                    f.write(resp.content)
                with Image.open(save_path) as img:
                    img.verify()
                return idx, save_path, "downloaded"
            else:
                return idx, None, f"HTTP {resp.status_code}"
        except Exception as e:
            if save_path.exists():
                try:
                    save_path.unlink()
                except Exception:
                    pass
            return idx, None, str(e)

    done_count = 0
    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        futures = {executor.submit(_fetch_one, t): t for t in tasks}
        for f in as_completed(futures):
            idx, p, st = f.result()
            results[idx] = (p, st)
            done_count += 1
            if progress_callback:
                progress_callback(done_count, total)

    return results


def run_audit_pipeline(
    excel_file_bytes: bytes,
    excel_filename: str,
    model_choice: str,
    output_base_dir: Path,
    status_callback=None,
    progress_callback=None
):
    """
    Executes the full end-to-end POG audit.
    Returns:
        dict with:
            - summary stats
            - detailed row dataframe list
            - path to audited excel
            - path to annotated images directory
            - zip buffer of annotated images
    """
    model_cfg = AUDIT_MODELS[model_choice]
    model = get_yolo_model(model_choice)

    timestamp = time.strftime("%Y%m%d_%H%M%S")
    clean_stem = "".join(c if c.isalnum() or c in "._-" else "_" for c in Path(excel_filename).stem)
    run_dir = output_base_dir / f"Audit_{model_cfg['id'].capitalize()}_{clean_stem}_{timestamp}"
    run_dir.mkdir(parents=True, exist_ok=True)

    download_dir = run_dir / "downloaded_images"
    download_dir.mkdir(parents=True, exist_ok=True)

    vis_dir = run_dir / "annotated_images"
    vis_dir.mkdir(parents=True, exist_ok=True)

    if status_callback:
        status_callback("Step 1 of 4: Parsing Excel workbook and mapping columns...")

    wb = openpyxl.load_workbook(io.BytesIO(excel_file_bytes), data_only=False)
    sheet_name = "Reset Photo Report" if "Reset Photo Report" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]

    all_rows = list(ws.iter_rows())
    if len(all_rows) < 2:
        raise ValueError("Workbook is empty or contains no data rows.")

    header_cells = all_rows[0]
    headers = [c.value for c in header_cells]
    photo_col_idx, store_col_idx, bay_col_idx = identify_columns(headers)

    total_data_rows = len(all_rows) - 1

    download_tasks = []
    for idx, row in enumerate(all_rows[1:], start=1):
        cell_photo = row[photo_col_idx]
        url = cell_photo.hyperlink.target if cell_photo.hyperlink else None
        photo_val = str(cell_photo.value or "").strip()
        if not url and photo_val.startswith(("http://", "https://")):
            url = photo_val
        store_val = str(row[store_col_idx].value or "").strip() if store_col_idx != -1 else f"Store_{idx}"
        download_tasks.append((idx, idx + 1, store_val, photo_val, url))

    # Phase 1: Image Fetching
    if status_callback:
        status_callback(f"Step 2 of 4: High-speed parallel download of {total_data_rows} bay photos...")

    def _dl_progress(done, total):
        if progress_callback:
            progress_callback(0.1 + (done / total) * 0.4)

    download_results = download_all_photos_parallel(
        download_tasks,
        download_dir,
        max_workers=20,
        progress_callback=_dl_progress
    )

    # Phase 2: YOLO Inference
    if status_callback:
        status_callback(f"Step 3 of 4: Running YOLO AI Vision ({model_cfg['name']}, Conf={model_cfg['conf']})...")

    row_results = []
    store_summary = {}
    approved_total = 0
    rejected_total = 0
    missing_image_count = 0

    import cv2

    for idx, row in enumerate(all_rows[1:], start=1):
        store_val = str(row[store_col_idx].value or "").strip() if store_col_idx != -1 else f"Store_{idx}"
        bay_val = str(row[bay_col_idx].value or "").strip() if bay_col_idx != -1 else ""

        img_path, dl_status = download_results.get(idx, (None, "Download failed"))
        vis_save_path = None

        if img_path and Path(img_path).exists():
            t0 = time.time()
            res = model.predict(
                str(img_path),
                conf=model_cfg["conf"],
                imgsz=640,
                device="cpu",
                verbose=False
            )[0]
            lat_ms = (time.time() - t0) * 1000.0

            boxes = res.boxes
            has_det = (boxes is not None and len(boxes) > 0)
            verdict = "Approved" if has_det else "Rejected"
            max_c = float(boxes.conf[0].cpu().numpy()) if has_det else 0.0

            classes_found = []
            if has_det:
                for b in boxes:
                    c_id = int(b.cls[0].cpu().numpy())
                    classes_found.append(model.names.get(c_id, f"class_{c_id}"))

            fmt_str = ", ".join(sorted(set(classes_found))) if classes_found else "none"
            box_count = len(boxes) if boxes is not None else 0

            # Always save annotated image for the web inspector
            plotted = res.plot()
            vis_file = vis_dir / f"audit_row_{idx+1}_{Path(img_path).name}"
            cv2.imwrite(str(vis_file), plotted)
            vis_save_path = str(vis_file)

        else:
            verdict = f"Error ({dl_status})"
            has_det = False
            max_c = 0.0
            fmt_str = "no_image"
            box_count = 0
            lat_ms = 0.0
            missing_image_count += 1

        if verdict == "Approved":
            approved_total += 1
        elif verdict.startswith("Error"):
            pass
        else:
            rejected_total += 1

        st_data = store_summary.setdefault(store_val, {"total": 0, "approved": 0, "rejected": 0, "error": 0})
        st_data["total"] += 1
        if verdict == "Approved":
            st_data["approved"] += 1
        elif verdict.startswith("Error"):
            st_data["error"] += 1
        else:
            st_data["rejected"] += 1

        row_results.append({
            "row_idx": idx + 1,
            "store": store_val,
            "bay": bay_val,
            "verdict": verdict,
            "confidence": round(max_c, 3),
            "format": fmt_str,
            "boxes": box_count,
            "latency_ms": round(lat_ms, 1),
            "original_path": str(img_path) if img_path else "",
            "annotated_path": vis_save_path or ""
        })

        if progress_callback:
            progress_callback(0.5 + (idx / total_data_rows) * 0.4)

    # Phase 3: Format Audited Excel Report
    if status_callback:
        status_callback("Step 4 of 4: Generating MSI Executive Dashboard Excel report...")

    header_fill = PatternFill(start_color="8B6F4E", end_color="8B6F4E", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    kpi_hdr_fill = PatternFill(start_color="212529", end_color="212529", fill_type="solid")
    pass_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    pass_font = Font(name="Calibri", size=10, bold=True, color="006100")
    fail_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    fail_font = Font(name="Calibri", size=10, bold=True, color="9C0006")
    err_fill = PatternFill(start_color="EFEFEF", end_color="EFEFEF", fill_type="solid")
    err_font = Font(name="Calibri", size=10, bold=True, color="555555")
    reg_font = Font(name="Calibri", size=10)
    bold_font = Font(name="Calibri", size=10, bold=True)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
    )

    start_col = len(headers) + 1
    new_cols = ["AI Audit Verdict", "AI Confidence", "AI Detected Type", "AI Detections Count"]
    for i, col_name in enumerate(new_cols):
        c = ws.cell(row=1, column=start_col + i, value=col_name)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")

    for res in row_results:
        r_num = res["row_idx"]
        c_verdict = ws.cell(row=r_num, column=start_col, value=res["verdict"])
        c_conf = ws.cell(row=r_num, column=start_col + 1, value=res["confidence"])
        c_fmt = ws.cell(row=r_num, column=start_col + 2, value=res["format"])
        c_box = ws.cell(row=r_num, column=start_col + 3, value=res["boxes"])

        for cell in (c_verdict, c_conf, c_fmt, c_box):
            cell.font = reg_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center")

        if res["verdict"] == "Approved":
            c_verdict.fill = pass_fill
            c_verdict.font = pass_font
        elif "Error" in res["verdict"]:
            c_verdict.fill = err_fill
            c_verdict.font = err_font
        else:
            c_verdict.fill = fail_fill
            c_verdict.font = fail_font

    # Executive Dashboard Sheet
    summary_sheet_title = "Audit Dashboard & Summary"
    if summary_sheet_title in wb.sheetnames:
        del wb[summary_sheet_title]
    ws_dash = wb.create_sheet(title=summary_sheet_title, index=0)

    ws_dash.cell(row=1, column=1, value="MSI SURFACES - PLANOGRAM (POG) AI AUDIT EXECUTIVE DASHBOARD").font = Font(name="Calibri", size=14, bold=True, color="8B6F4E")

    rate_pct = (approved_total / total_data_rows * 100.0) if total_data_rows > 0 else 0.0
    stores_approved = sum(1 for s in store_summary.values() if s["approved"] >= 1)
    stores_review = len(store_summary) - stores_approved
    store_pass_rate = (stores_approved / len(store_summary) * 100.0) if len(store_summary) > 0 else 0.0

    kpis = [
        ("Audit Target Specification", model_cfg["name"]),
        ("Model Weights Architecture", model_cfg["weights"].name),
        ("Calibrated Confidence Threshold", f"{model_cfg['conf']:.3f}"),
        ("Total Photo Bays Audited", total_data_rows),
        ("Total Compliant Photos (Approved)", approved_total),
        ("Total Non-Compliant Photos (Rejected)", rejected_total),
        ("Overall Photo Pass Rate (%)", f"{rate_pct:.2f}%"),
        ("Total Store Locations Audited", len(store_summary)),
        ("Compliant Store Locations (>=1 Passed Bay)", stores_approved),
        ("Store Locations Requiring Review (0 Passed Bays)", stores_review),
        ("Store Compliance Pass Rate (%)", f"{store_pass_rate:.2f}%"),
        ("Missing / Unreachable Image URLs", missing_image_count)
    ]

    ws_dash.cell(row=3, column=1, value="Executive KPI Metric").fill = kpi_hdr_fill
    ws_dash.cell(row=3, column=1).font = header_font
    ws_dash.cell(row=3, column=2, value="Audit Outcome Value").fill = kpi_hdr_fill
    ws_dash.cell(row=3, column=2).font = header_font

    for idx, (k, v) in enumerate(kpis, start=4):
        c1 = ws_dash.cell(row=idx, column=1, value=k)
        c2 = ws_dash.cell(row=idx, column=2, value=v)
        c1.font = bold_font
        c2.font = reg_font
        c1.border = thin_border
        c2.border = thin_border
        if "Approved" in k or "Compliant Store" in k:
            c2.fill = pass_fill
            c2.font = pass_font
        elif "Rejected" in k or "Requiring Review" in k:
            c2.fill = fail_fill
            c2.font = fail_font

    # Store Breakdown in Excel
    st_start_row = len(kpis) + 6
    ws_dash.cell(row=st_start_row, column=1, value="STORE LOCATION AUDIT BREAKDOWN").font = Font(name="Calibri", size=12, bold=True, color="8B6F4E")

    st_headers = ["Store / Location", "Total Bays Audited", "Approved", "Rejected", "Errors", "Store Pass Rate (%)", "Compliance Status"]
    for col_idx, h_text in enumerate(st_headers, start=1):
        c = ws_dash.cell(row=st_start_row + 1, column=col_idx, value=h_text)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center")

    def sort_store(s):
        try:
            return (0, int(s))
        except ValueError:
            return (1, str(s))

    for idx, st_num in enumerate(sorted(store_summary.keys(), key=sort_store), start=st_start_row + 2):
        s_data = store_summary[st_num]
        s_rate = (s_data["approved"] / s_data["total"] * 100.0) if s_data["total"] > 0 else 0.0
        # If at least 1 photo for a store passes detection, the entire store is Approved
        s_status = "Approved" if s_data["approved"] >= 1 else "Needs Review"

        vals = [
            st_num,
            s_data["total"],
            s_data["approved"],
            s_data["rejected"],
            s_data["error"],
            f"{s_rate:.1f}%",
            s_status
        ]
        for c_idx, val in enumerate(vals, start=1):
            cell = ws_dash.cell(row=idx, column=c_idx, value=val)
            cell.font = reg_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center")
            if c_idx == 7:
                if s_status == "Approved":
                    cell.fill = pass_fill
                    cell.font = pass_font
                else:
                    cell.fill = fail_fill
                    cell.font = fail_font

    for col in ws_dash.columns:
        max_l = max(len(str(c.value or "")) for c in col)
        col_letter = get_column_letter(col[0].column)
        ws_dash.column_dimensions[col_letter].width = max(max_l + 3, 14)

    output_excel_path = run_dir / f"{Path(excel_filename).stem}_Audited.xlsx"
    wb.save(output_excel_path)
    wb.close()

    # Create ZIP buffer for annotated images
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for img_file in vis_dir.glob("*.jpg"):
            zf.write(img_file, arcname=img_file.name)
        for img_file in vis_dir.glob("*.png"):
            zf.write(img_file, arcname=img_file.name)
    zip_buffer.seek(0)

    if progress_callback:
        progress_callback(1.0)

    return {
        "run_dir": str(run_dir),
        "total_rows": total_data_rows,
        "approved_count": approved_total,
        "rejected_count": rejected_total,
        "error_count": missing_image_count,
        "pass_rate": rate_pct,
        "model_name": model_cfg["name"],
        "confidence_threshold": model_cfg["conf"],
        "store_summary": store_summary,
        "row_results": row_results,
        "output_excel_path": str(output_excel_path),
        "zip_buffer": zip_buffer.getvalue()
    }
